"""Independent real-source fixture for testing the Vega renderer (not Desktop)."""
from pathlib import Path
from openpyxl import load_workbook
import csv,json,runpy,collections,math
ROOT=Path(__file__).resolve().parents[1];BASE=ROOT/'powerbi-ledningsrapport';DATA=BASE/'data'
source=runpy.run_path(str(ROOT/'scripts/budget-scenario-data.py'));actual=source['actuals']()
read=lambda name:list(csv.DictReader((DATA/(name+'.csv')).open(encoding='utf-8-sig')))
nodes=read('Budgetrapportnoder');mapping=read('Budgetrader');budget=collections.defaultdict(dict)
wb=load_workbook(DATA/'Budgetscenarier.xlsx',data_only=True)
scenarios={}
for scenario,sheet in [(1,'2026_Fastställd'),(2,'2026_Försiktigt'),(3,'2026_Offensivt')]:
 data=collections.defaultdict(dict)
 for row in wb[sheet].iter_rows(min_row=2,values_only=True):
  for m,v in enumerate(row[4:16],1):data[202600+m][row[2]]=v
 scenarios[scenario]=data
budget=scenarios[1]
wb.close()
def raw(dataset,months):
 total=collections.defaultdict(int)
 for m in months:
  for k,v in dataset[202600+m].items():total[k]+=v
 return total
def value(dataset,node,months):
 vals=raw(dataset,months);ids=[int(r['KontoID']) for r in mapping if r['Rad']==node['Rad'] and (not node['Undergrupp'] or r['Undergrupp']==node['Undergrupp']) and (not node['Post'] or r['Post']==node['Post'])]
 if len(ids)==1 and ids[0]<0:
  sales=sum(v for k,v in vals.items() if k<2000);gp=sales+sum(v for k,v in vals.items() if 2000<=k<3000);ebit=sum(vals.values())
  return {-1:gp,-2:gp/sales,-3:ebit,-4:ebit/sales}[ids[0]]
 return sum(vals[k] for k in ids)
def make(months,scenario=1,comparison=0):
 result=[]
 for n in nodes:
  for m in months+[13]:
   period=months if m==13 else [m];a=value(scenarios[scenario],n,period);b=value(scenarios[comparison] if comparison else actual,n,period)
   for slot in ([1,2,3] if m==13 else [1,2]):
    margin=n['Typ']=='marginal';v=a if slot==1 else b if slot==2 else (a-b if comparison else b-a)*(100 if margin else 1)
    if not margin:v/=1000
    label=(f'{v:+.1f} pp' if slot==3 else f'{v*100:.1f} %').replace('.',',') if margin else (f'{v:+,.0f}' if slot==3 else f'{v:,.0f}').replace(',','\u00a0')
    variance=None
    if slot==3:
     difference=a-b if comparison else b-a;reference=b if comparison else a
     variance=difference*100 if margin else difference/abs(reference) if reference else None
     rounded=round(variance if margin else variance*100,1) if variance is not None else 0
     label='—' if variance is None else (format(rounded,'+.1f' if rounded else '.1f')+(' pp' if margin else ' %')).replace('.',',').replace('-','−')
    labels=['A · '+source['SCENARIOS'][scenario],'B · '+source['SCENARIOS'][comparison],'Avvikelse'] if comparison else ['Budget','Utfall','Avvikelse']
    result.append(dict(Node=n['NodID'],Label=n['Etikett'],Parent=n['Förälder'] or None,Ancestor=n['Överförälder'] or None,Depth=int(n['Nivå']),Kind=n['Typ'],Order=int(n['Ordning']),Children=int(n['HarBarn']),Section=n['Avsnitt'] or None,Group='Vald period' if m==13 else source['MONTHS'][m-1],GroupOrder=m,Metric=labels[slot-1],Slot=slot,Amount=v,Text=label,Color=('#B42318' if rounded<0 else '#007956' if rounded>0 else '#5B6B80') if slot==3 else '#5B6B80' if margin else '#0F3163'))
 return result
standard=make([7,8,9]);wide=make(list(range(1,10)))
live=json.loads((BASE/'validation/budget-live-result.json').read_text(encoding='utf-8-sig'))['resultSets'][1]
root={ (r['Label'],r['Group'],r['Metric']):r['Amount'] for r in standard if r['Depth']==0 }
assert len(root)==81
for r in live:assert math.isclose(root[r['Budgetrader[Rad]'],r['Budgetkolumner[Grupp]'],r['Budgetkolumner[Kolumn]']],r['[Värde]'],abs_tol=1e-7)
(BASE/'validation/budget-statement-fixture.json').write_text(json.dumps(standard,ensure_ascii=False),encoding='utf-8')
(BASE/'validation/budget-statement-wide-fixture.json').write_text(json.dumps(wide,ensure_ascii=False),encoding='utf-8')
(BASE/'validation/budget-statement-scenarios-fixture.json').write_text(json.dumps(make(list(range(1,13)),2,1),ensure_ascii=False),encoding='utf-8')
print('Fixture root amounts match all 81 previously executed Desktop DAX cells.')
