"""Vary demo budget detail while preserving every monthly RR group total."""
from pathlib import Path
import json
import runpy
from collections import defaultdict
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
source=runpy.run_path(str(root/'scripts/budget-scenario-data.py'))
path=root/'powerbi-ledningsrapport/data/Budgetscenarier.xlsx'
wb=load_workbook(path)
marker='Detaljvariation v1: huvudgruppernas månadsbudgetar är oförändrade.'
info=wb.worksheets[0]
if any(row[0]==marker for row in info.iter_rows(values_only=True)):
    print('Detail variation already applied; workbook unchanged.')
    raise SystemExit(0)
groups=defaultdict(list)
for account,group,*_ in source['ACCOUNTS']:
    groups[group].append(account)
checks=0
for ws in wb.worksheets[1:]:
    rows={int(ws.cell(r,3).value):r for r in range(2,ws.max_row+1)}
    for month in range(1,13):
        for ids in groups.values():
            if len(ids)<2:
                continue
            old=[ws.cell(rows[k],month+4).value for k in ids]
            total=sum(old)
            # Repeatable account/month profiles; no random refresh drift.
            weights=[abs(v)*(1+(((k*7+month*3)%13)-6)*.025) for k,v in zip(ids,old)]
            denom=sum(weights)
            if not denom:
                continue
            new=[round(total*w/denom) for w in weights]
            new[max(range(len(new)),key=lambda i:weights[i])]+=total-sum(new)
            assert sum(new)==total
            assert all(v==0 or (v>0)==(total>0) for v in new)
            for k,v in zip(ids,new):
                ws.cell(rows[k],month+4).value=v
            checks+=1
info.append([marker])
info.append(['Underkategorier har olika konto- och månadsprofiler. Utfall och huvudgruppernas totaler har inte ändrats.'])
wb.save(path)
report={'status':'passed','preservedMonthlyGroupTotals':checks,'actualsChanged':False,'method':'deterministic redistribution within each existing RR group'}
(root/'powerbi-ledningsrapport/validation/budget-detail-variation.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report))
