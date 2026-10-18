"""Create the editable Excel budget and additive RR mapping. Never rewrite actuals.
Run --reset-workbook only to deliberately replace an edited demo workbook.
"""
from pathlib import Path
import csv, json, collections, sys, hashlib
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'powerbi-ledningsrapport'; DATA=BASE/'data'
MONTHS=['Januari','Februari','Mars','April','Maj','Juni','Juli','Augusti','September','Oktober','November','December']
SCENARIOS={1:'Huvudscenario',2:'Försiktigt scenario',3:'Offensivt scenario'}
# ID, main RR row, subrow, leaf, main order, suborder, leaforder
ACCOUNTS=[
 (1001,'Nettoomsättning','Produktförsäljning','',1,1,0),(1002,'Nettoomsättning','Teknisk service','',1,2,0),(1003,'Nettoomsättning','Bemanning','',1,3,0),
 (2001,'Direkta kostnader','Varor och material','',2,1,0),(2002,'Direkta kostnader','Direkt leveranspersonal','',2,2,0),(2003,'Direkta kostnader','Underleverantörer','',2,3,0),(2004,'Direkta kostnader','Övrigt direkt','',2,4,0),(2005,'Direkta kostnader','Ej uppdragsfördelad personal','',2,5,0),
 (3001,'Gemensam personal','Löner och sociala avgifter','',5,1,0),
 (4001,'Övriga rörelsekostnader','Lokaler','',6,1,0),
 (4002,'Övriga rörelsekostnader','IT','Licenser',6,2,1),(4003,'Övriga rörelsekostnader','IT','IT-drift',6,2,2),(4004,'Övriga rörelsekostnader','IT','IT-konsulter',6,2,3),
 (4005,'Övriga rörelsekostnader','Marknadsföring','',6,3,0),(4006,'Övriga rörelsekostnader','Administration och övrigt','',6,4,0),
 (5001,'Avskrivningar','Avskrivningar enligt plan','',7,1,0)]
CALC=[(-1,'Bruttovinst','','',3,0,0),(-2,'Bruttomarginal','','',4,0,0),(-3,'Rörelseresultat','','',8,0,0),(-4,'Rörelsemarginal','','',9,0,0)]
def rows(name):return list(csv.DictReader((DATA/(name+'.csv')).open(encoding='utf-8-sig')))
def write(name,columns,data):
 with (DATA/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.writer(f);w.writerow(columns);w.writerows(data)
def rounded(v):return int(Decimal(str(v)).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
def actuals():
 values=collections.defaultdict(lambda:collections.defaultdict(int))
 for r in rows('Försäljning'):values[int(r['Datumnyckel'])//100][1000+int(r['Områdesnyckel'])]+=int(r['Nettointäkt'])
 dm={'Varukostnad':2001,'Direkt leveranspersonal':2002,'Underleverantörer':2003,'Övrigt direkt':2004,'Ej uppdragsfördelad personal':2005}
 for r in rows('Direkta kostnader'):values[int(r['Datumnyckel'])//100][dm[r['Kostnadstyp']]]-=int(r['Belopp'])
 for r in rows('Gemensamma kostnader'):
  dest=values[int(r['Datumnyckel'])//100];v=int(r['Belopp'])
  if r['Kostnadstyp']=='Administration':
   parts={k:rounded(v*f) for k,f in [(3001,.6),(4002,.10),(4003,.06),(4004,.04),(4005,.08)]}
   parts[4006]=v-sum(parts.values())
   for k,n in parts.items():dest[k]-=n
  else:dest[4001 if r['Kostnadstyp']=='Lokaler och drift' else 5001]-=v
 return values
def main():
 actual=actuals(); targets={int(r['Datumnyckel'])//100:int(r['Rörelseresultat']) for r in rows('Budget')}
 source_hash={n:hashlib.sha256((DATA/(n+'.csv')).read_bytes()).hexdigest() for n in ['Försäljning','Direkta kostnader','Gemensamma kostnader','Budget','Uppdrag','Kunder','Artiklar']}
 path=DATA/'Budgetscenarier.xlsx'
 if not path.exists() or '--reset-workbook' in sys.argv:
  wb=Workbook();info=wb.active;info.title='Läs mig'
  for row in [ ['Budgetscenarier – redigerbar demobudget'],['Belopp anges i kronor. Ändra blå beloppsceller och spara filen innan Uppdatera i Power BI.'],['Tre flikar per år 2026. Historiken har enbart Huvudscenario; inga tomma alternativ erbjuds.'],['Endast kontorader importeras. Inga Excel-summor ingår i faktatabellen.'],['Huvudscenario är avstämd mot ursprunglig Budget.csv per månad (EBIT).'],['Historisk detaljbudget: omsättning 95 % av utfall, direkta kostnader 94 % av utfall.'],['Gemensamma kostnader fördelas proportionellt för att nå befintligt budgeterat EBIT.'],['Försiktigt scenario: omsättning −8 %, direkta kostnader −5 %, gemensamma kostnader +2 %.'],['Offensivt scenario: omsättning +10 %, direkta kostnader +6 %, gemensamma kostnader +3 %.'],['Avskrivningar oförändrade mellan scenarier. Inga formler kräver Excel-omräkning.'],['November/december 2026 bygger på samma månader 2025 med 10 % tillväxt.'],['Ändra sökvägen via Power Query-parametern BudgetWorkbook. Utfall ändras aldrig här.'] ]:info.append(row)
  info.column_dimensions['A'].width=125
  for year in [2023,2024,2025,2026]:
   for scenario in ([1,2,3] if year==2026 else [1]):
    ws=wb.create_sheet(f'{year}_{["Fastställd","Försiktigt","Offensivt"][scenario-1]}')
    ws.append(['År','ScenarioID','KontoID','Post']+MONTHS)
    monthvalues={}
    for month in range(1,13):
     key=year*100+month; src=actual.get(key,actual.get((year-1)*100+month,{}));growth=1 if key in actual else 1.1
     vals={k:rounded(src.get(k,0)*growth*(.95 if k<2000 else .94 if k<3000 else 1)) for k,*_ in ACCOUNTS}
     ebit=targets.get(key,rounded(targets.get((year-1)*100+month,0)*1.1))
     costkeys=[k for k,*_ in ACCOUNTS if k>=3000]
     remainder=sum(v for k,v in vals.items() if k<3000)-ebit
     denom=-sum(vals[k] for k in costkeys)
     for k in costkeys[:-1]:vals[k]=-rounded(remainder*(-vals[k])/denom)
     vals[costkeys[-1]]=-remainder-sum(vals[k] for k in costkeys[:-1])
     if scenario!=1:
      for k in vals:
       factor=(.92 if k<2000 else .95 if k<3000 else 1 if k>=5000 else 1.02) if scenario==2 else (1.10 if k<2000 else 1.06 if k<3000 else 1 if k>=5000 else 1.03)
       vals[k]=rounded(vals[k]*factor)
     monthvalues[month]=vals
    for k,group,sub,leaf,*_ in ACCOUNTS:ws.append([year,scenario,k,leaf or sub]+[monthvalues[m][k] for m in range(1,13)])
    table=Table(displayName=f'Budget_{year}_{scenario}',ref=f'A1:P{len(ACCOUNTS)+1}');table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    ws.freeze_panes='E2';ws.column_dimensions['D'].width=35
    for col in ['A','B','C']:ws.column_dimensions[col].width=12
    for row in ws.iter_rows(min_row=2,min_col=5):
     for c in row:c.number_format='#,##0;[Red]-#,##0';c.font=Font(name='Calibri',color='0F3163');c.fill=PatternFill('solid',fgColor='EDF5FF')
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='0F3163')
    for col in 'EFGHIJKLMNOP':ws.column_dimensions[col].width=16
   
  wb.save(path)
  # Newly generated demo detail must not repeat one variance across all accounts.
  import runpy
  runpy.run_path(str(ROOT/'scripts/vary-budget-details.py'))
 # Utility dimensions and mapping are reproducible; workbook is preserved by default.
 write('Budgetkonton',['KontoID','RRgrupp','Undergrupp','Post','Gruppordning','Underordning','Postordning'],ACCOUNTS)
 write('Budgetrader',['KontoID','Rad','Undergrupp','Post','Radordning','Underordning','Postordning'],ACCOUNTS+CALC)
 write('Budgetscenarion',['ScenarioID','Scenario','Ordning'],[(i,n,i) for i,n in SCENARIOS.items()])
 write('Budgetjämförelsescenario',['ScenarioID','Scenario','Ordning'],[(i,n,i) for i,n in SCENARIOS.items()])
 write('Budgetår',['År'],[(y,) for y in [2024,2025,2026]])
 write('Budgetmånader',['Nummer','Månad','Kortnamn'],[(m,MONTHS[m-1],MONTHS[m-1][:3]) for m in range(1,13)])
 write('Budgetperioder',['Månadsnyckel','År','Nummer'],[(y*100+m,y,m) for y in range(2023,2027) for m in range(1,13)])
 write('Budgetjämförelse',['Läge','Ordning'],[('Scenario mot utfall',1),('Jämför scenarier',2)])
 write('Budgetvisning',['Visning','Ordning'],[('Månad',1),('Ackumulerat',2)])
 columns=[]
 for mode in [1,2]:
  for a in SCENARIOS:
   for b in ([0] if mode==1 else [i for i in SCENARIOS if i!=a]):
    for m in range(1,14):
     for slot in ([1,2,3] if m==13 else [1,2]):
      label=('Budget' if slot==1 else 'Utfall') if mode==1 else (f'A · {SCENARIOS[a]}' if slot==1 else f'B · {SCENARIOS[b]}')
      if slot==3:label='Avvikelse'
      columns.append([mode,a,b,m,'Vald period' if m==13 else MONTHS[m-1],slot,label])
 write('Budgetkolumner',['LägeID','ScenarioA','ScenarioB','MånadID','Grupp','Position','Kolumn'],columns)
 manifest={'workbook':str(path.relative_to(ROOT)),'sourceActualHashes':source_hash,'latestCompleteActualMonth':202609,'historicalScenarios':'Huvudscenario only 2023–2025','assumptions':{'adminToPersonnel':.6,'adminToLicenses':.10,'adminToITOperations':.06,'adminToITConsultants':.04,'adminToMarketing':.08,'adminResidual':.12},'referenceActualQ3':{str(k):sum(actual[202600+m][k] for m in [7,8,9]) for k,*_ in ACCOUNTS}}
 (BASE/'validation/budget-source-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 print('Budget workbook ready; existing actual and baseline CSV files unchanged.')
if __name__=='__main__':main()
