"""Independent workbook/CSV and PBIR checks. Does not claim live DAX/UI verification."""
from pathlib import Path
import sys,json,csv,hashlib,math,collections,os
from openpyxl import load_workbook
import runpy
source=runpy.run_path(str(Path(__file__).with_name('budget-scenario-data.py')))
ROOT=source['ROOT'];BASE=source['BASE'];DATA=source['DATA'];actual=source['actuals']();accounts=source['ACCOUNTS']
workbook=DATA/'Budgetscenarier.xlsx'
def read_budget():
 wb=load_workbook(workbook,data_only=True);result=collections.defaultdict(dict)
 for ws in wb.worksheets:
  if not ws.title[:4].isdigit():continue
  for row in ws.iter_rows(min_row=2,values_only=True):
   year,scenario,account,_=row[:4]
   for month,value in enumerate(row[4:16],1):
    assert isinstance(value,(int,float)) and value is not None
    key=(year*100+month,scenario)
    assert account not in result[key]
    result[key][account]=value
 wb.close();return result
budget=read_budget()
expected_accounts={x[0] for x in accounts}
assert all(set(v)==expected_accounts for v in budget.values())
assert len(budget)==72  # 36 historical baseline months + 36 scenario-months for 2026.
manifest=json.loads((BASE/'validation/budget-source-manifest.json').read_text(encoding='utf-8'))
for name,expected in manifest['sourceActualHashes'].items():assert hashlib.sha256((DATA/(name+'.csv')).read_bytes()).hexdigest()==expected,name
targets={int(r['Datumnyckel'])//100:int(r['Rörelseresultat']) for r in source['rows']('Budget')}
for key,target in targets.items():assert sum(budget[key,1].values())==target,(key,target)
assert sum(sum(budget[key,1].values()) for key in range(202510,202513))+sum(sum(budget[key,1].values()) for key in range(202601,202610))==9750000
def summary(values):
 revenue=sum(v for k,v in values.items() if k<2000);direct=sum(v for k,v in values.items() if 2000<=k<3000)
 gp=revenue+direct;ebit=sum(values.values());return {'revenue':revenue,'grossProfit':gp,'EBIT':ebit,'grossMargin':gp/revenue if revenue else None,'operatingMargin':ebit/revenue if revenue else None}
def aggregate(dataset,year,months,scenario=None):
 vals=collections.defaultdict(int)
 for m in months:
  row=dataset.get((year*100+m,scenario),{}) if scenario else dataset.get(year*100+m,{})
  for k,v in row.items():vals[k]+=v
 return vals
r12=collections.defaultdict(int)
for k,vals in actual.items():
 if 202510<=k<=202609:
  for account,value in vals.items():r12[account]+=value
assert [summary(r12)[n] for n in ['revenue','grossProfit','EBIT']]==[139500000,36040000,8000000]
reference=summary(aggregate(actual,2026,[7,8,9]));reference_budget=summary(aggregate(budget,2026,[7,8,9],1))
cases=[]
for year,scenarios in [(2024,[1]),(2025,[1]),(2026,[1,2,3])]:
 for scenario in scenarios:
  for months in [[7],[7,8,9],[1,3,9],list(range(1,10 if year==2026 else 13))]:
   a=aggregate(actual,year,months);b=aggregate(budget,year,months,scenario)
   # Every common cost is allocated exactly once; direct delivery personnel remain direct.
   for m in months:
    key=year*100+m
    common=sum(int(r['Belopp']) for r in source['rows']('Gemensamma kostnader') if int(r['Datumnyckel'])//100==key)
    assert -sum(v for k,v in actual[key].items() if k>=3000)==common
   accumulated=[summary(aggregate(budget,year,months[:i+1],scenario)) for i in range(len(months))]
   assert accumulated[-1]==summary(b)
   assert summary(a)['EBIT']==sum(a.values())
   cases.append({'year':year,'scenario':scenario,'months':months,'actual':summary(a),'budget':summary(b)})
for a in [1,2,3]:
 for b in [1,2,3]:
  if a==b:continue
  va=summary(aggregate(budget,2026,list(range(1,13)),a));vb=summary(aggregate(budget,2026,list(range(1,13)),b))
  cases.append({'scenarioA':a,'scenarioB':b,'year':2026,'months':'full year incl. future months','A':va,'B':vb,'differenceEBIT':va['EBIT']-vb['EBIT']})
# Read-back test: Excel is the actual editable source, not a generated CSV replacement.
# Restore the exact original bytes even if the test fails.
original=workbook.read_bytes()
try:
 wb=load_workbook(workbook);ws=wb['2026_Fastställd'];before=ws.cell(2,11).value;ws.cell(2,11,before+1000);wb.save(workbook);wb.close()
 changed=read_budget();assert changed[202607,1][1001]==budget[202607,1][1001]+1000
 assert summary(aggregate(changed,2026,[7,8,9],1))['EBIT']==reference_budget['EBIT']+1000
finally:workbook.write_bytes(original)
assert read_budget()==budget
report=BASE/os.environ.get('PBI_BUDGET_REPORT','LedningsrapportUtanScroll.Report');metadata=json.loads((report/'definition/pages/pages.json').read_text(encoding='utf-8-sig'))
assert len(metadata['pageOrder'])==(7 if (BASE/'Ledningsrapport.SemanticModel/definition/tables/Planinställning.tmdl').exists() else 5)
page_id=json.loads((BASE/'validation/budget-page-manifest.json').read_text(encoding='utf-8'))['page']
page=report/'definition/pages'/page_id
visuals=[json.loads(p.read_text(encoding='utf-8-sig')) for p in page.glob('visuals/*/visual.json')]
assert any(v['visual']['visualType']=='image' and 'Aubeco' in json.dumps(v) for v in visuals)
assert not any('Områdesval' in json.dumps(v) for v in visuals)
for v in visuals:
 p=v['position'];assert 0<=p['x'] and 0<=p['y'] and p['x']+p['width']<=1920 and p['y']+p['height']<=1200
matrix=next(v for v in visuals if v['visual']['visualType']=='deneb7E15AEF80B9E4D4F8E12924291ECE89A')
assert not any(v['visual']['visualType']=='pivotTable' for v in visuals)
assert len(matrix['visual']['query']['queryState']['dataset']['projections'])==16
spec=json.loads(matrix['visual']['objects']['vega'][0]['properties']['jsonSpec']['expr']['Literal']['Value'][1:-1].replace("''", "'"))
assert any(d['name']=='open' for d in spec['data'])
assert any(m['name']=='body' and m['clip'] for m in spec['marks'])
assert matrix['position']['height']==572
for kind in ['expand','collapse','reset']:
 name=hashlib.sha256(('budget-scenarios:'+kind).encode()).hexdigest()[:20]
 bookmark=json.loads((report/'definition/bookmarks'/(name+'.bookmark.json')).read_text(encoding='utf-8'))
 assert matrix['name'] in bookmark['options']['targetVisualNames']
result={'status':'passed','scope':'Independent persisted workbook/CSV calculations and PBIR configuration; not live DAX or UI','baselineMonthlyReconciliations':len(targets),'cases':cases,'referenceQ3Actual':reference,'referenceQ3Budget':reference_budget,'excelEditReadRestoreTest':True,'actualSourcesUnchanged':True,'defaultExpansionConfigured':True,'UIInteractionsVerified':False,'liveExcelRefreshVerified':False,'screenshotsVerified':False}
(BASE/'validation/budget-checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:result[k] for k in ['status','scope','baselineMonthlyReconciliations','referenceQ3Actual','referenceQ3Budget','excelEditReadRestoreTest']},ensure_ascii=False))
